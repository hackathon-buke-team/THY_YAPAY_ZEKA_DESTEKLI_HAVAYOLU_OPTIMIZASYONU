# -*- coding: utf-8 -*-
"""Dokuman 2 — Problemin Boyutu ve Cozum Altyapisi."""
import os
import sys
import importlib

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from builder import Doc, BLUE_D, INK, INK2, MUTED

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\darkb\Desktop\Buke_cozum\docs\Buke_PROBLEM_BOYUTU_ve_COZUCU.docx"

MODULES = ["c2_bolum1", "c2_bolum2", "c2_ek"]


def cover(D):
    d = D.d
    D.spacer(14)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("TEKNOFEST · Yapay Zekâ Destekli Havayolu Optimizasyonu")
    r.font.size = Pt(11); r.font.color.rgb = MUTED; r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("PROBLEMİN BOYUTU")
    r.font.size = Pt(30); r.bold = True; r.font.color.rgb = BLUE_D; r.font.name = "Segoe UI"
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("ve ÇÖZÜM ALTYAPISI")
    r.font.size = Pt(30); r.bold = True; r.font.color.rgb = BLUE_D; r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Kaç ihtimal var · neyi değiştiriyoruz · Pyomo ve çözücü ne işe yarıyor?")
    r.font.size = Pt(12.5); r.font.color.rgb = INK2; r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Takım: Büke")
    r.font.size = Pt(12); r.bold = True; r.font.color.rgb = INK; r.font.name = "Segoe UI"

    D.img("g3_arama_uzayi.png", None, 16.2)

    D.spacer(4)
    D.tbl(["", "Değer"],
          [["Ham girdi verisi", "**89.094** satır · 4 Excel dosyası"],
           ["Türetilmiş veri katmanı", "**118.177** satır · 9 CSV"],
           ["Değiştirdiğimiz karar", "**4.553** uçuş olayının IST saati (±60 dk / 10 dk ızgara)"],
           ["Olay başına seçenek", "**13**"],
           ["Tarife karar uzayı", "**13⁴⁵⁵³ ≈ 10⁵·⁰⁷²** ihtimal"],
           ["Model boyutu", "**100.566** ikili değişken · **70.359** kısıt"],
           ["Modelleme çerçevesi", "**Pyomo** (problemi tarif eder)"],
           ["Çözücü", "**HiGHS** (problemi çözer ve optimalite kanıtı üretir)"],
           ["Sonuç", "**%+4,72** rekabetçilik · global MIP-gap **%1,18**"]],
          [5.0, 12.0], 9.2)

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Tüm sayılar projenin gerçek çıktı ve veri dosyalarından hesaplanmıştır.")
    r.font.size = Pt(9); r.italic = True; r.font.color.rgb = MUTED; r.font.name = "Segoe UI"


def toc(D):
    D.pagebreak()
    D.h2("İçindekiler")
    rows = [
        ["**1**", "Bizden tam olarak ne isteniyor?",
         "değiştirdiğimiz tek şey · dokunmadıklarımız · amaç fonksiyonu"],
        ["**2**", "Elimizde kaç adet veri var?",
         "4 girdi dosyası · derived/ katmanı · output/ · veri hunisi"],
        ["**3**", "Kaç ihtimal var?",
         "13⁴⁵⁵³ hesabı · karşılaştırmalar · uzayı neyle küçültüyoruz"],
        ["**4**", "Pyomo nedir, ne işe yarar?",
         "modelleme ↔ çözme ayrımı · üç işi · ne DEĞİLDİR"],
        ["**5**", "Çözücü nedir, nasıl çalışır?",
         "LP gevşemesi · dallanma-sınırlama · kesme · sezgisel · MIP-gap"],
        ["**6**", "Hangi çözücü, neden?",
         "Gurobi · CPLEX · HiGHS · CBC karşılaştırması"],
        ["**7**", "`appsi` katmanı ve warm-start",
         "kalıcı arayüz · setSolution · 08 neden buna bağımlı"],
        ["**8**", "Modelin gerçek boyutu",
         "100.566 ikili değişken dökümü · kısıt sayıları · sonuç"],
        ["**A**", "Ek A — tek sayfalık sayı künyesi", "tüm rakamlar tek tabloda"],
        ["**B**", "Ek B — iki konunun adım adım anlatımı",
         "823 sabit olay neden kımıldamıyor · huni figürü satır satır"],
    ]
    D.tbl(["#", "Bölüm", "İçerik"], rows, [1.2, 6.8, 9.0], 9.2)


def main():
    sys.path.insert(0, HERE)
    D = Doc()
    cover(D)
    toc(D)
    for mod in MODULES:
        m = importlib.import_module(mod)
        print("[render]", mod, f"({len(m.ITEMS)} oge)")
        D.render(m.ITEMS)
    D.save(OUT, "Problemin Boyutu ve Çözüm Altyapısı")


if __name__ == "__main__":
    main()
