# -*- coding: utf-8 -*-
"""ALGORITMA — SIFIRDAN ANLATIM · DOCX kurucu.

Calistirma:  python docs/kod_rehberi_kaynak/main4.py
Cikti:       docs/Buke_ALGORITMA_SIFIRDAN_ANLATIM.docx
"""
import os
import sys
import importlib

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from builder import Doc, BLUE, BLUE_D, INK, INK2, MUTED, GREEN, RED

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "Buke_ALGORITMA_SIFIRDAN_ANLATIM.docx")

MODULES = ["c4_bolum1", "c4_bolum2", "c4_bolum3"]


class Doc4(Doc):
    """builder.Doc + 'eq' (ortalanmis formul satiri)."""

    def eq(self, text):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(9)
        self._rich(p, text, 11.5, BLUE_D)

    def render(self, items):
        for kind, val in items:
            if kind == "eq":
                self.eq(val)
            else:
                super().render([(kind, val)])


def cover(D):
    d = D.d
    for _ in range(3):
        D.spacer(10)

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("TEKNOFEST · Yapay Zekâ Destekli Havayolu Optimizasyonu")
    r.font.size = Pt(11); r.font.color.rgb = MUTED; r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("ALGORİTMA")
    r.font.size = Pt(40); r.bold = True; r.font.color.rgb = BLUE_D; r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("sıfırdan, adım adım — jüri sunumu hazırlığı")
    r.font.size = Pt(14); r.font.color.rgb = INK2; r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Takım: Büke")
    r.font.size = Pt(12); r.bold = True; r.font.color.rgb = INK; r.font.name = "Segoe UI"

    D.img("f01_pipeline.png", None, 15.5)

    D.spacer(8)
    D.box("Bu doküman üç soruya cevap vermek için yazıldı",
          "**1)** \"E2 ne, anlamadım\"  →  Bölüm 6.   "
          "**2)** \"MIP gap %1,18 kesin mi, nereden biliyoruz?\"  →  Bölüm 9.   "
          "**3)** \"100.552 ikili değişken ne demek?\"  →  Bölüm 4 ve Ek A.\n"
          "Ayrıca projenin tamamı — ham veriden nihai sonuca — baştan sona anlatılıyor.",
          "info")

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("Tüm sayılar kaynak koddan ve koşu raporlarından doğrulanmıştır; "
                  "ikili değişken dökümü yeniden sayılarak teyit edilmiştir.")
    r.font.size = Pt(8.6); r.italic = True; r.font.color.rgb = MUTED; r.font.name = "Segoe UI"


def toc(D):
    D.pagebreak()
    D.h2("İçindekiler")
    rows = [
        ["**0**", "Bu doküman ne işe yarar?", "60 saniyelik özet · okuma sırası"],
        ["**1**", "Problem — neye karar veriyoruz?",
         "aktarma ağı · J = E1 + boşluk + E2 · rekabetçilik"],
        ["**2**", "Veri — dört Excel dosyasından modele",
         "sağlık kontrolü · sarma tuzağı · ρ · ayarlanabilir/sabit"],
        ["**3**", "Optimizasyon 101 — sıfırdan, jargonsuz",
         "ikili değişken · LP gevşetmesi · dallanma-sınırlama · incumbent/bound"],
        ["**4**", "Modelin değişkenleri — yüz binin dökümü",
         "x · y · z · w · β · γ  ·  ön-işleme"],
        ["**5**", "Kısıtlar — kurallar ve gerekçeleri",
         "A1–G · büyük-M · B2 çatışması · ceza terimi"],
        ["**6**", "**E2 — sorduğunuz sorunun tam cevabı**",
         "**iki farklı E2 · min'in doğrusallaştırılması · neden kapalı**"],
        ["**7**", "Amaç fonksiyonu", "iki ödül, bir ceza · hangi sayıyı raporluyoruz"],
        ["**8**", "Çözüm hattı", "güne göre ayrıştırma · sıcak başlangıç · çift bariyer"],
        ["**9**", "**MIP gap — \"%1,18 kesin mi?\"**",
         "**matematiksel kaynak · kodun 941. satırı · neyi garanti eder, neyi etmez**"],
        ["**10**", "Bağımsız doğrulama", "çift kayıt · yakalanan üç gerçek hata"],
        ["**11**", "Sonuçlar", "ana tablo · gün kırılımı · duyarlılık"],
        ["**12**", "Jüri soru bankası", "10 soru, hazır cevaplar"],
        ["**A**", "Ek A — Sayı künyesi ve sayım kanıtı",
         "iki modelin karşılaştırması · dosya haritası · bilinen tutarsızlıklar"],
    ]
    D.tbl(["", "Bölüm", "İçerik"], rows, [1.2, 6.4, 9.4], 8.8)


def main():
    D = Doc4()
    cover(D)
    toc(D)
    for name in MODULES:
        mod = importlib.import_module(name)
        importlib.reload(mod)
        D.render(mod.ITEMS)
    D.save(OUT, footer_title="Algoritma · Sıfırdan Anlatım")


if __name__ == "__main__":
    sys.path.insert(0, HERE)
    os.chdir(HERE)
    main()
