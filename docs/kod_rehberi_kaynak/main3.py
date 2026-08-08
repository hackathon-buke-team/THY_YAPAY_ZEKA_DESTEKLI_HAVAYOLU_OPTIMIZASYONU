# -*- coding: utf-8 -*-
"""Dokuman 3 — Ozgunluk, Terimler ve Juri Soru Bankasi."""
import os
import sys
import importlib

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from builder import Doc, BLUE_D, INK, INK2, MUTED

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\darkb\Desktop\Buke_cozum\docs\Buke_OZGUNLUK_TERIMLER_JURI_SORULARI.docx"

MODULES = ["c3_ozgunluk", "c3_terimler", "c3_sorular"]


def cover(D):
    d = D.d
    D.spacer(12)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("TEKNOFEST · Yapay Zekâ Destekli Havayolu Optimizasyonu")
    r.font.size = Pt(11); r.font.color.rgb = MUTED; r.font.name = "Segoe UI"

    for txt, sz in [("ÖZGÜNLÜK, TERİMLER", 27), ("ve JÜRİ SORULARI", 27)]:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10 if sz == 27 and txt.startswith("ÖZ") else 0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(txt)
        r.font.size = Pt(sz); r.bold = True; r.font.color.rgb = BLUE_D
        r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Biz ne yaptık · her terim sıfırdan · 45 soru ve cevabı")
    r.font.size = Pt(12.5); r.font.color.rgb = INK2; r.font.name = "Segoe UI"

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Takım: Büke")
    r.font.size = Pt(12); r.bold = True; r.font.color.rgb = INK; r.font.name = "Segoe UI"

    D.img("h1_katki_haritasi.png", None, 14.6)

    D.spacer(2)
    D.tbl(["Bölüm", "İçerik"],
          [["**1 · Biz ne yaptık?**",
            "\"Model yaptık mı?\" — kelimenin iki anlamı · çözücü ne yapar, biz ne yaptık"],
           ["**2 · Özgünlük**",
            "Formülasyon · çözüm stratejisi · veri mühendisliği · doğrulama · raporlama"],
           ["**3 · Yapmasaydık ne olurdu?**",
            "10 tasarım kararının ölçülmüş alternatifi"],
           ["**4 · Neler bizim değil**", "Dürüst sınır çizimi"],
           ["**5 · Terimler**",
            "MIP · big-M · one-hot · warm-start · MIP-gap · λ cezası · 40+ terim, sıfırdan"],
           ["**6 · Jüri soru bankası**",
            "45 soru · problem · model · çözücü · veri · doğrulama · eleştiri"]],
          [4.6, 12.4], 9.0)

    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("Tüm sayılar projenin gerçek çıktı ve veri dosyalarından hesaplanmıştır.")
    r.font.size = Pt(9); r.italic = True; r.font.color.rgb = MUTED; r.font.name = "Segoe UI"


def main():
    sys.path.insert(0, HERE)
    D = Doc()
    cover(D)
    for mod in MODULES:
        m = importlib.import_module(mod)
        print("[render]", mod, f"({len(m.ITEMS)} oge)")
        D.render(m.ITEMS)
    D.save(OUT, "Özgünlük, Terimler ve Jüri Soruları")


if __name__ == "__main__":
    main()
