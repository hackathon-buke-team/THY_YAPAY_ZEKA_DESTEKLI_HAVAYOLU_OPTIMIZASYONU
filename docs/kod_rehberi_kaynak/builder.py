# -*- coding: utf-8 -*-
"""DOCX ureticisi: icerik listesini (content.py) Word dosyasina cevirir.

Kod bloklari GERCEK kaynak dosyalardan satir araligiyla okunur -> birebir dogru.
"""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
SRCROOT = r"C:\Users\darkb\Desktop\Buke_cozum"

# ---------------------------------------------------------------- palet
INK = RGBColor(0x14, 0x14, 0x14)
INK2 = RGBColor(0x45, 0x44, 0x42)
MUTED = RGBColor(0x7d, 0x7b, 0x76)
BLUE = RGBColor(0x1c, 0x5d, 0xa8)
BLUE_D = RGBColor(0x14, 0x42, 0x78)
RED = RGBColor(0xb8, 0x33, 0x32)
GREEN = RGBColor(0x1f, 0x6b, 0x52)
KW = RGBColor(0x0b, 0x57, 0xa4)
STR_ = RGBColor(0x9a, 0x5b, 0x14)
CMT = RGBColor(0x62, 0x77, 0x62)
NUM = RGBColor(0x14, 0x6b, 0x62)
CODE_BG = "f4f3ef"
BOX_BG = "eef3fa"
WARN_BG = "fbf1ee"
OK_BG = "eef6f1"
HDR_BG = "e8e7e1"

KEYWORDS = {
    "False", "None", "True", "and", "as", "assert", "async", "await", "break", "class",
    "continue", "def", "del", "elif", "else", "except", "finally", "for", "from", "global",
    "if", "import", "in", "is", "lambda", "nonlocal", "not", "or", "pass", "raise", "return",
    "try", "while", "with", "yield",
}
TOKEN_RE = re.compile(r"([A-Za-z_][A-Za-z_0-9]*|\d+\.?\d*|.)", re.S)


# ---------------------------------------------------------------- xml yardimcilari
def _shd(el, fill):
    pPr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_par(par, fill):
    _shd(par._p, fill)


def par_border(par, sides=("left",), color="c9c8c2", sz=18, space=6):
    pPr = par._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    for s in sides:
        e = OxmlElement("w:" + s)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), str(space))
        e.set(qn("w:color"), color)
        pbdr.append(e)


def keep_with_next(par, val=True):
    pPr = par._p.get_or_add_pPr()
    e = OxmlElement("w:keepNext")
    if not val:
        e.set(qn("w:val"), "0")
    pPr.append(e)


def cell_shade(cell, fill):
    _shd(cell._tc.get_or_add_tcPr(), fill)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    trPr.append(e)


def add_field(par, instr):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "…"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2); r._r.append(t); r._r.append(f3)
    return r


# ---------------------------------------------------------------- sozdizimi renklendirme
def tokenize_line(line, state):
    """(parcalar, yeni_state) — parca = (metin, tip). tip: txt|kw|str|cmt|num"""
    out = []
    i = 0
    n = len(line)
    if state:                                    # ucul tirnak icindeyiz
        idx = line.find(state)
        if idx == -1:
            return [(line, "str")], state
        out.append((line[:idx + 3], "str"))
        i = idx + 3
        state = None
    buf = ""
    while i < n:
        ch = line[i]
        three = line[i:i + 3]
        if three in ('"""', "'''"):
            if buf:
                out.extend(_words(buf)); buf = ""
            end = line.find(three, i + 3)
            if end == -1:
                out.append((line[i:], "str"))
                return out, three
            out.append((line[i:end + 3], "str"))
            i = end + 3
            continue
        if ch in "\"'":
            if buf:
                out.extend(_words(buf)); buf = ""
            j = i + 1
            while j < n:
                if line[j] == "\\":
                    j += 2; continue
                if line[j] == ch:
                    break
                j += 1
            out.append((line[i:j + 1], "str"))
            i = j + 1
            continue
        if ch == "#":
            if buf:
                out.extend(_words(buf)); buf = ""
            out.append((line[i:], "cmt"))
            return out, None
        buf += ch
        i += 1
    if buf:
        out.extend(_words(buf))
    return out, None


def _words(s):
    res = []
    for tok in TOKEN_RE.findall(s):
        if tok in KEYWORDS:
            res.append((tok, "kw"))
        elif tok and tok[0].isdigit():
            res.append((tok, "num"))
        else:
            res.append((tok, "txt"))
    # ardisik ayni tipleri birlestir
    merged = []
    for t, k in res:
        if merged and merged[-1][1] == k:
            merged[-1] = (merged[-1][0] + t, k)
        else:
            merged.append([t, k]) if False else merged.append((t, k))
    out = []
    for t, k in merged:
        if out and out[-1][1] == k:
            out[-1] = (out[-1][0] + t, k)
        else:
            out.append((t, k))
    return out


COLOR = {"kw": KW, "str": STR_, "cmt": CMT, "num": NUM, "txt": INK}


# ---------------------------------------------------------------- belge kurucu
class Doc:
    def __init__(self):
        self.d = Document()
        self._setup()
        self._src_cache = {}

    # ---------------- sayfa & stiller
    def _setup(self):
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(1.8)
        st = self.d.styles["Normal"]
        st.font.name = "Segoe UI"
        st.font.size = Pt(10)
        st.font.color.rgb = INK
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
        pf = st.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.22

    def footer(self, title="Kod Rehberi"):
        for sec in self.d.sections:
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"Büke · THY Tarife Optimizasyonu — {title}   ·   ")
            r.font.size = Pt(7.5); r.font.color.rgb = MUTED; r.font.name = "Segoe UI"
            add_field(p, " PAGE ")
            for rr in p.runs:
                rr.font.size = Pt(7.5); rr.font.color.rgb = MUTED; rr.font.name = "Segoe UI"

    # ---------------- kaynak okuma
    def src(self, fname):
        if fname not in self._src_cache:
            path = os.path.join(SRCROOT, "src", fname) if fname != "run_all.py" \
                else os.path.join(SRCROOT, fname)
            with open(path, encoding="utf-8") as f:
                self._src_cache[fname] = f.read().split("\n")
        return self._src_cache[fname]

    # ---------------- ogeler
    def h1(self, text, num=None):
        # NOT: add_page_break() ayri bir paragraf ekler; sayfa tam dolmussa o bos
        # paragraf sonraki sayfaya taşınır ve ARDINDAN kırar -> bos sayfa. Bunun
        # yerine basligin kendisine "oncesinde sayfa kir" ozelligi verilir.
        p = self.d.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if num:
            r = p.add_run(num + "  ")
            r.font.size = Pt(21); r.bold = True; r.font.color.rgb = RGBColor(0xc8, 0xc7, 0xc1)
            r.font.name = "Segoe UI"
        r = p.add_run(text)
        r.font.size = Pt(19); r.bold = True; r.font.color.rgb = BLUE_D; r.font.name = "Segoe UI"
        line = self.d.add_paragraph()
        line.paragraph_format.space_before = Pt(0)
        line.paragraph_format.space_after = Pt(11)
        par_border(line, sides=("bottom",), color="1c5da8", sz=10, space=2)

    def h2(self, text):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(15)
        p.paragraph_format.space_after = Pt(5)
        keep_with_next(p)
        r = p.add_run(text)
        r.font.size = Pt(13.5); r.bold = True; r.font.color.rgb = BLUE_D; r.font.name = "Segoe UI"

    def h3(self, text):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        keep_with_next(p)
        r = p.add_run(text)
        r.font.size = Pt(11); r.bold = True; r.font.color.rgb = INK; r.font.name = "Segoe UI"

    def p(self, text, size=10, color=INK, italic=False, after=6, before=0, align=None):
        par = self.d.add_paragraph()
        par.paragraph_format.space_after = Pt(after)
        par.paragraph_format.space_before = Pt(before)
        if align == "c":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._rich(par, text, size, color, italic)
        return par

    def _rich(self, par, text, size=10, color=INK, italic=False):
        """**kalin**, `kod`, *italik* destegi."""
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = par.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(size - 1.2)
                r.font.color.rgb = RGBColor(0x33, 0x44, 0x66)
                continue
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = par.add_run(part[1:-1]); r.italic = True
            else:
                r = par.add_run(part)
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = "Segoe UI"
            if italic:
                r.italic = True

    def bul(self, items, size=9.8):
        for it in items:
            par = self.d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.62)
            par.paragraph_format.first_line_indent = Cm(-0.32)
            par.paragraph_format.space_after = Pt(3)
            r = par.add_run("▪  ")
            r.font.size = Pt(size); r.font.color.rgb = BLUE; r.font.name = "Segoe UI"
            self._rich(par, it, size)

    def num(self, items, size=9.8):
        for i, it in enumerate(items, 1):
            par = self.d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.72)
            par.paragraph_format.first_line_indent = Cm(-0.42)
            par.paragraph_format.space_after = Pt(3)
            r = par.add_run(f"{i}.  ")
            r.font.size = Pt(size); r.bold = True; r.font.color.rgb = BLUE; r.font.name = "Segoe UI"
            self._rich(par, it, size)

    # ---------------- kod blogu
    def code(self, fname, a, b, note=None):
        lines = self.src(fname)
        cap = self.d.add_paragraph()
        cap.paragraph_format.space_before = Pt(9)
        cap.paragraph_format.space_after = Pt(1)
        keep_with_next(cap)
        r = cap.add_run(f"src/{fname}" if fname != "run_all.py" else "run_all.py")
        r.font.size = Pt(7.8); r.bold = True; r.font.color.rgb = BLUE; r.font.name = "Consolas"
        r = cap.add_run(f"   ·   satır {a}–{b}" + (f"   ·   {note}" if note else ""))
        r.font.size = Pt(7.8); r.font.color.rgb = MUTED; r.font.name = "Segoe UI"

        # blok dosyanin ortasindan basliyorsa (ör. docstring icinden), giris
        # durumunu bastan tarayarak bul -> ucul tirnak icindeki apostroflar
        # yanlislikla string baslangici sayilmasin
        state = None
        for prev in lines[:a - 1]:
            _, state = tokenize_line(prev.rstrip("\n"), state)
        first, last = a - 1, b
        for i in range(first, last):
            raw = lines[i].rstrip("\n")
            par = self.d.add_paragraph()
            pf = par.paragraph_format
            pf.space_before = Pt(0); pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.12)
            shade_par(par, CODE_BG)
            par_border(par, sides=("left",), color="b9c6d8", sz=18, space=5)
            ln = par.add_run(f"{i + 1:>4}  ")
            ln.font.name = "Consolas"; ln.font.size = Pt(6.5)
            ln.font.color.rgb = RGBColor(0xa8, 0xa7, 0xa1)
            parts, state = tokenize_line(raw, state)
            if not parts:
                parts = [(" ", "txt")]
            for txt, kind in parts:
                run = par.add_run(txt.replace("\t", "    "))
                run.font.name = "Consolas"
                run.font.size = Pt(7.6)
                run.font.color.rgb = COLOR[kind]
                if kind == "kw":
                    run.bold = True
                if kind == "cmt":
                    run.italic = True
        # blok sonu bosluk
        tail = self.d.add_paragraph()
        tail.paragraph_format.space_before = Pt(0)
        tail.paragraph_format.space_after = Pt(0)
        tail.paragraph_format.line_spacing = 0.5
        r = tail.add_run(" ")
        r.font.size = Pt(3)

    # ---------------- kutu
    def box(self, title, text, kind="info"):
        fill = {"info": BOX_BG, "warn": WARN_BG, "ok": OK_BG}[kind]
        col = {"info": BLUE_D, "warn": RED, "ok": GREEN}[kind]
        bc = {"info": "9ab6dc", "warn": "e0b6b0", "ok": "aacfbf"}[kind]
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        cell = t.cell(0, 0)
        cell.width = Cm(17.0)
        cell_shade(cell, fill)
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        if title:
            p0 = cell.add_paragraph()
            p0.paragraph_format.space_after = Pt(2)
            p0.paragraph_format.space_before = Pt(3)
            r = p0.add_run(title)
            r.bold = True; r.font.size = Pt(9.6); r.font.color.rgb = col; r.font.name = "Segoe UI"
        p1 = cell.add_paragraph()
        p1.paragraph_format.space_after = Pt(4)
        if not title:
            p1.paragraph_format.space_before = Pt(4)
        self._rich(p1, text, 9.3, INK2)
        self._tbl_borders(t, bc)
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(3)
        r = sp.add_run(" "); r.font.size = Pt(3)

    def _tbl_borders(self, t, color="c9c8c2", sz=6, inner=True):
        tbl = t._tbl
        tblPr = tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        names = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inner else [])
        for nm in names:
            e = OxmlElement("w:" + nm)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
            borders.append(e)
        tblPr.append(borders)

    # ---------------- tablo
    def tbl(self, headers, rows, widths=None, size=8.6, hdr_size=8.6):
        t = self.d.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        hdr = t.rows[0]
        set_repeat_header(hdr)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            cell_shade(c, HDR_BG)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(h)
            r.bold = True; r.font.size = Pt(hdr_size); r.font.color.rgb = INK
            r.font.name = "Segoe UI"
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                p = cells[i].paragraphs[0]
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                self._rich(p, str(v), size, INK2)
                if ri % 2 == 1:
                    cell_shade(cells[i], "faf9f6")
        if widths:
            for r_ in t.rows:
                for i, w in enumerate(widths):
                    r_.cells[i].width = Cm(w)
        self._tbl_borders(t)
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
        r = sp.add_run(" "); r.font.size = Pt(4)

    # ---------------- gorsel
    def img(self, name, caption=None, width=16.4):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        path = os.path.join(IMG, name) if not os.path.isabs(name) else name
        p.add_run().add_picture(path, width=Cm(width))
        if caption:
            c = self.d.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_after = Pt(10)
            r = c.add_run(caption)
            r.font.size = Pt(8.2); r.italic = True; r.font.color.rgb = MUTED
            r.font.name = "Segoe UI"

    def qa(self, num, text):
        """Soru basligi — numarali, acik mavi zeminli."""
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        cell = t.cell(0, 0)
        cell.width = Cm(17.0)
        cell_shade(cell, "e7eefa")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        keep_with_next(p)
        r = p.add_run(f"S{num}  ")
        r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = BLUE
        r.font.name = "Segoe UI"
        self._rich(p, text, 10.2, BLUE_D)
        for run in p.runs[1:]:
            run.bold = True
        self._tbl_borders(t, "9ab6dc")
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(2)
        r = sp.add_run(" "); r.font.size = Pt(2)

    def pagebreak(self):
        self.d.add_page_break()

    def spacer(self, pt=8):
        p = self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(" "); r.font.size = Pt(pt)

    # ---------------- render
    def render(self, items):
        for kind, val in items:
            if kind == "h1":
                self.h1(val[1], val[0]) if isinstance(val, tuple) else self.h1(val)
            elif kind == "h2":
                self.h2(val)
            elif kind == "h3":
                self.h3(val)
            elif kind == "p":
                self.p(val)
            elif kind == "pm":                       # muted / kucuk not
                self.p(val, size=9.2, color=MUTED, italic=True)
            elif kind == "bul":
                self.bul(val)
            elif kind == "num":
                self.num(val)
            elif kind == "code":
                self.code(*val)
            elif kind == "tbl":
                self.tbl(*val)
            elif kind == "img":
                self.img(*val)
            elif kind == "box":
                self.box(val[0], val[1], val[2] if len(val) > 2 else "info")
            elif kind == "qa":
                self.qa(val[0], val[1])
            elif kind == "pb":
                self.pagebreak()
            elif kind == "sp":
                self.spacer(val or 8)
            else:
                raise ValueError("bilinmeyen oge: " + kind)

    def save(self, path, footer_title="Kod Rehberi"):
        self.footer(footer_title)
        self.d.save(path)
        print("[docx]", path, f"{os.path.getsize(path)//1024} KB")
