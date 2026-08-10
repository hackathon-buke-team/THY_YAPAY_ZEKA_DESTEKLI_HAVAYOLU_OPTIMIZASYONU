# -*- coding: utf-8 -*-
"""Sunum Çalışma Kitabı — her slaydın görseli + o slaytta anlatılacaklar.

Sunumdaki 10 slaydı PNG olarak dışa aktarır, konuşmacı notlarıyla birlikte tek
bir markdown'a yazar ve teslim hattındaki gibi .docx / .pdf üretir.

Çalıştırma:  python src/12_sunum_calisma_kitabi.py [sunum.pptx]
Üretir:      docs/figures/sunum/slaytNN.png
             docs/Buke_SUNUM_CALISMA_KITABI.{md,docx,pdf}
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGDIR = os.path.join(ROOT, "docs", "figures", "sunum")
MD = os.path.join(ROOT, "docs", "Buke_SUNUM_CALISMA_KITABI.md")

PPTX = sys.argv[1] if len(sys.argv) > 1 else "Buke_Juri_Sunumu_v3_SADE3.pptx"
PPTX = PPTX if os.path.isabs(PPTX) else os.path.join(ROOT, PPTX)

PAGEBREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'


def export_slides(pptx):
    """PowerPoint COM ile slaytları PNG'ye aktarır. Dosya açıksa kopya üzerinden çalışır."""
    os.makedirs(FIGDIR, exist_ok=True)
    tmp = os.path.join(tempfile.gettempdir(), "buke_deck_export.pptx")
    shutil.copy(pptx, tmp)
    ps = (
        "$ErrorActionPreference='Stop';"
        "$pp = New-Object -ComObject PowerPoint.Application;"
        f"$pres = $pp.Presentations.Open('{tmp}', $true, $false, $false);"
        "for ($i=1; $i -le $pres.Slides.Count; $i++) {"
        f"  $pres.Slides($i).Export(('{FIGDIR}\\slayt{{0:00}}.png' -f $i), 'PNG', 1920, 1080) }};"
        "$pres.Close()"
    )
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                       capture_output=True, text=True)
    if r.returncode != 0:
        print(f"[uyari] slayt dışa aktarımı başarısız, mevcut PNG'ler kullanılacak\n{r.stderr.strip()}")


def clean(note):
    """Konuşmacı notunu markdown'a uygun hale getirir; SÜRE satırını ayırır."""
    note = note.replace("\r", "").strip()
    sure = ""
    m = re.match(r"^SÜRE:\s*([^\n]+)\n?", note)
    if m:
        sure = m.group(1).strip()
        note = note[m.end():].strip()
    bloklar = [b.strip() for b in re.split(r"\n\s*\n", note) if b.strip()]
    return sure, bloklar


def slayt_basligi(slide, i):
    """İç sayfalarda başlık sol üstteki metin kutusudur; kapak/kapanışta en büyük punto."""
    from pptx.util import Inches

    metinler = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        tam = " ".join(x.strip() for x in sh.text_frame.text.strip().split("\n") if x.strip())
        if len(tam) < 4 or re.fullmatch(r"\d+\s*/\s*10", tam):
            continue
        punto = max((r.font.size.pt for pa in sh.text_frame.paragraphs
                     for r in pa.runs if r.font.size), default=0)
        metinler.append((sh.top or 0, sh.left or 0, punto, tam))

    ust = [m for m in metinler if m[0] < Inches(0.8) and m[1] < Inches(4)]
    if ust:
        return min(ust)[3]                        # şablonun sol üst başlık kutusu
    serit_disi = [m for m in metinler if not (m[3] == m[3].upper() and len(m[3]) > 20)]
    return max(serit_disi or metinler, key=lambda m: m[2])[3] if metinler else f"Slayt {i}"


def main():
    export_slides(PPTX)
    prs = Presentation(PPTX)

    L = ["# Sunum Çalışma Kitabı", "",
         "**Takım Büke** · TEKNOFEST — Yapay Zekâ Destekli Havayolu Optimizasyonu", "",
         "Her sayfada bir slayt ve o slaytta anlatılacaklar var. Toplam sunum süresi "
         "10 dakika, ardından 5 dakika soru-cevap.", "",
         f"*Kaynak sunum: `{os.path.basename(PPTX)}`*", ""]

    toplam = []
    for i, s in enumerate(prs.slides, 1):
        png = f"figures/sunum/slayt{i:02d}.png"
        note = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        sure, bloklar = clean(note)
        if sure:
            toplam.append(sure)

        baslik = slayt_basligi(s, i)

        L += [PAGEBREAK if i > 1 else "", f"## {i}. {baslik}", ""]
        if sure:
            L += [f"**Süre:** {sure}", ""]
        L += [f"![]({png})", "", "### Ne anlatacağız", ""]
        L += [b + "\n" for b in bloklar] or ["*(not yok)*\n"]

    L += [PAGEBREAK, "## Süre planı", "",
          "| Slayt | Süre |", "|---|---|"]
    for i, s in enumerate(toplam, 1):
        L.append(f"| {i} | {s} |")
    L += ["", "Sunum 10 dakika ile sınırlı. Provada her slaytta kalan süreyi kontrol edin; "
          "5. ve 9. slaytlar en kritik olanlar, gerekirse 7. ve 8. slaytları hızlandırın.", ""]

    with open(MD, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
    print(f"[md  ] {os.path.relpath(MD, ROOT)}")

    try:
        import pypandoc
        docx = MD.replace(".md", ".docx")
        pypandoc.convert_file(MD, "docx", outputfile=docx,
                              extra_args=["--resource-path", os.path.join(ROOT, "docs")])
        print(f"[docx] {os.path.relpath(docx, ROOT)}")
        pdf = docx.replace(".docx", ".pdf")
        ps = ("$ErrorActionPreference='Stop';"
              "$w = New-Object -ComObject Word.Application; $w.Visible=$false;"
              f"$d = $w.Documents.Open('{docx}'); $d.SaveAs('{pdf}', 17);"
              "$d.Close($false); $w.Quit()")
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                           capture_output=True, text=True)
        print(f"[pdf ] {os.path.relpath(pdf, ROOT)}" if os.path.exists(pdf)
              else f"[pdf ] üretilemedi: {r.stderr.strip()[:200]}")
    except Exception as e:
        print(f"[atla] docx/pdf üretilemedi: {e}")


main()
